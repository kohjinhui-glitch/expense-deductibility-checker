"""
Convert Deductibility.xlsx → data.json for the expense deductibility web app.

Reads the "Deductibility" sheet and writes a JSON array of row objects, one per
data row. Column headers become object keys; empty cells become empty strings.
Run at Docker build time — the web app loads data.json directly, no SheetJS needed.
"""

import json
import sys
import openpyxl

XLSX_PATH = "Deductibility.xlsx"
SHEET_NAME = "Deductibility"
JSON_PATH = "data.json"


def convert():
    try:
        wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    except FileNotFoundError:
        sys.exit(f"ERROR: {XLSX_PATH} not found")

    if SHEET_NAME not in wb.sheetnames:
        available = ", ".join(wb.sheetnames)
        sys.exit(f'ERROR: Sheet "{SHEET_NAME}" not found. Available: {available}')

    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        sys.exit(f'ERROR: Sheet "{SHEET_NAME}" is empty')

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    data = []
    for row in rows[1:]:
        obj = {}
        for header, value in zip(headers, row):
            if header == "":
                continue
            if value is None:
                obj[header] = ""
            else:
                obj[header] = str(value).strip()
        # Skip rows where every value is empty or "Please select"
        values = list(obj.values())
        if all(v == "" or v == "Please select" for v in values):
            continue
        data.append(obj)

    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    print(f"Converted {len(data)} rows → {JSON_PATH}")


if __name__ == "__main__":
    convert()
